import hashlib
import io
import json
import os
import re
import secrets
import sqlite3
import unicodedata
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path

import pandas as pd
from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

try:
    from google import genai
except ImportError:
    genai = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
DB_PATH = os.getenv("DATABASE_PATH", str(BASE_DIR / "assessment.db"))
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
ALLOWED_NOTES = {"pdf", "docx"}
ALLOWED_LOGOS = {"png", "jpg", "jpeg", "webp"}
ALLOWED_IMAGES = {"png", "jpg", "jpeg", "webp", "gif"}
EXAM_TYPES = ("midterm", "final", "quiz", "practice", "mock", "assignment", "other")
DATA_CLEAR_PASSWORD = os.getenv("DATA_CLEAR_PASSWORD", "change-this-clear-password")
THEMES = {"ocean", "forest", "royal", "sunset", "slate"}
app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "change-this-development-secret")


def db():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    return connection


def query(sql, params=(), one=False):
    connection = db()
    rows = connection.execute(sql, params).fetchall()
    connection.close()
    return (rows[0] if rows else None) if one else rows


def execute(sql, params=()):
    connection = db()
    cursor = connection.execute(sql, params)
    connection.commit()
    last_id = cursor.lastrowid
    connection.close()
    return last_id


def hash_password(password):
    salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 120_000).hex()
    return f"{salt}${digest}"


def check_password(password, stored):
    try:
        salt, digest = stored.split("$", 1)
        candidate = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 120_000).hex()
        return secrets.compare_digest(candidate, digest)
    except ValueError:
        return False


def init_db():
    connection = db()
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL,
            full_name TEXT NOT NULL, password_hash TEXT NOT NULL,
            role TEXT NOT NULL CHECK(role IN ('admin', 'teacher', 'student')),
            class_name TEXT NOT NULL DEFAULT '', subjects TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT, subject TEXT NOT NULL DEFAULT 'General',
            topic TEXT NOT NULL, class_name TEXT NOT NULL DEFAULT '', image_path TEXT NOT NULL DEFAULT '',
            stem TEXT NOT NULL, option_a TEXT NOT NULL,
            option_b TEXT NOT NULL, option_c TEXT NOT NULL, option_d TEXT NOT NULL,
            correct_option TEXT NOT NULL, difficulty TEXT NOT NULL DEFAULT 'medium',
            explanation TEXT, standard TEXT,
            created_by INTEGER, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS exams (
            id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL,
            subject TEXT NOT NULL DEFAULT 'General', class_name TEXT NOT NULL DEFAULT '',
            exam_type TEXT NOT NULL DEFAULT 'other', instructions TEXT,
            duration_minutes INTEGER NOT NULL DEFAULT 60, created_by INTEGER NOT NULL,
            published INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS exam_questions (
            exam_id INTEGER NOT NULL, question_id INTEGER NOT NULL,
            position INTEGER NOT NULL, PRIMARY KEY (exam_id, question_id)
        );
        CREATE TABLE IF NOT EXISTS attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT, exam_id INTEGER NOT NULL,
            student_id INTEGER NOT NULL, score REAL NOT NULL, total INTEGER NOT NULL,
            started_at TEXT NOT NULL, submitted_at TEXT NOT NULL, UNIQUE(exam_id, student_id)
        );
        CREATE TABLE IF NOT EXISTS responses (
            attempt_id INTEGER NOT NULL, question_id INTEGER NOT NULL,
            answer TEXT, correct INTEGER NOT NULL, PRIMARY KEY (attempt_id, question_id)
        );
        CREATE TABLE IF NOT EXISTS rewrite_permissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT, attempt_id INTEGER NOT NULL,
            granted_by INTEGER NOT NULL, granted_at TEXT NOT NULL,
            UNIQUE(attempt_id)
        );
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY, value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS subject_access_codes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, subject TEXT NOT NULL,
            code TEXT NOT NULL, expires_at TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1,
            created_by INTEGER NOT NULL, created_at TEXT NOT NULL
        );
        """
    )
    columns = {row["name"] for row in connection.execute("PRAGMA table_info(questions)").fetchall()}
    if "subject" not in columns:
        connection.execute("ALTER TABLE questions ADD COLUMN subject TEXT NOT NULL DEFAULT 'General'")
    if "difficulty" not in columns:
        connection.execute("ALTER TABLE questions ADD COLUMN difficulty TEXT NOT NULL DEFAULT 'medium'")
    if "class_name" not in columns:
        connection.execute("ALTER TABLE questions ADD COLUMN class_name TEXT NOT NULL DEFAULT ''")
    if "image_path" not in columns:
        connection.execute("ALTER TABLE questions ADD COLUMN image_path TEXT NOT NULL DEFAULT ''")
    user_columns = {row["name"] for row in connection.execute("PRAGMA table_info(users)").fetchall()}
    if "class_name" not in user_columns:
        try:
            connection.execute("ALTER TABLE users ADD COLUMN class_name TEXT NOT NULL DEFAULT ''")
        except sqlite3.OperationalError as error:
            if "duplicate column name" not in str(error).lower():
                raise
    if "subjects" not in user_columns:
        try:
            connection.execute("ALTER TABLE users ADD COLUMN subjects TEXT NOT NULL DEFAULT ''")
        except sqlite3.OperationalError as error:
            if "duplicate column name" not in str(error).lower():
                raise
    exam_columns = {row["name"] for row in connection.execute("PRAGMA table_info(exams)").fetchall()}
    attempt_columns = {row["name"] for row in connection.execute("PRAGMA table_info(attempts)").fetchall()}
    if "started_at" not in attempt_columns:
        connection.execute("ALTER TABLE attempts ADD COLUMN started_at TEXT NOT NULL DEFAULT ''")
    if "subject" not in exam_columns:
        connection.execute("ALTER TABLE exams ADD COLUMN subject TEXT NOT NULL DEFAULT 'General'")
    if "duration_minutes" not in exam_columns:
        connection.execute("ALTER TABLE exams ADD COLUMN duration_minutes INTEGER NOT NULL DEFAULT 60")
    if "exam_type" not in exam_columns:
        connection.execute("ALTER TABLE exams ADD COLUMN exam_type TEXT NOT NULL DEFAULT 'other'")
    if "class_name" not in exam_columns:
        try:
            connection.execute("ALTER TABLE exams ADD COLUMN class_name TEXT NOT NULL DEFAULT ''")
        except sqlite3.OperationalError as error:
            if "duplicate column name" not in str(error).lower():
                raise
    connection.execute(
        """UPDATE exams SET class_name = (
            SELECT q.class_name FROM exam_questions eq
            JOIN questions q ON q.id = eq.question_id
            WHERE eq.exam_id = exams.id AND q.class_name != ''
            ORDER BY eq.position LIMIT 1
        ) WHERE class_name = '' AND EXISTS (
            SELECT 1 FROM exam_questions eq2
            JOIN questions q2 ON q2.id = eq2.question_id
            WHERE eq2.exam_id = exams.id AND q2.class_name != ''
        )"""
    )
    for exam_row in connection.execute("SELECT id, class_name FROM exams WHERE class_name != ''").fetchall():
        normalized_class = class_group(exam_row["class_name"])
        if normalized_class != exam_row["class_name"]:
            connection.execute("UPDATE exams SET class_name = ? WHERE id = ?", (normalized_class, exam_row["id"]))
    if connection.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
        now = datetime.now().isoformat(timespec="seconds")
        connection.executemany(
            "            INSERT INTO users(username, full_name, password_hash, role, class_name, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            [
                ("admin", "System Administrator", hash_password("admin123"), "admin", "", now),
                ("teacher", "Demo Teacher", hash_password("teacher123"), "teacher", "", now),
                ("student", "Demo Student", hash_password("student123"), "student", "Demo Class", now),
            ],
        )
    connection.commit()
    connection.close()


def current_user():
    user_id = session.get("user_id")
    return query("SELECT * FROM users WHERE id = ? AND active = 1", (user_id,), one=True) if user_id else None


def teacher_subjects(user=None):
    user = user or current_user()
    if not user or user["role"] != "teacher":
        return []
    return [value.strip() for value in (user["subjects"] or "").split(",") if value.strip()]


def scoped_question_sql(user):
    if user["role"] == "teacher" and teacher_subjects(user):
        placeholders = ",".join("?" for _ in teacher_subjects(user))
        return f" AND (created_by = ? OR subject IN ({placeholders}))", [user["id"], *teacher_subjects(user)]
    if user["role"] == "teacher":
        return " AND created_by = ?", [user["id"]]
    return "", []


def login_required(*roles):
    def decorator(function):
        @wraps(function)
        def wrapped(*args, **kwargs):
            user = current_user()
            if not user:
                return redirect(url_for("login"))
            if roles and user["role"] not in roles:
                flash("You do not have permission to access that page.", "error")
                return redirect(url_for("dashboard"))
            return function(*args, **kwargs)
        return wrapped
    return decorator


def school_settings():
    values = {
        "school_name": "School Assessment Hub", "school_address": "", "school_logo": "",
        "theme": "ocean", "theme_image": "",
    }
    values.update({row["key"]: row["value"] for row in query("SELECT key, value FROM settings")})
    return values


def category_values(key, fallback_sql):
    configured = [value.strip() for value in school_settings().get(key, "").split(",") if value.strip()]
    if configured:
        return sorted(set(configured))
    return [row["value"] for row in query(fallback_sql)]


def extract_notes(file_storage):
    filename = secure_filename(file_storage.filename or "")
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in ALLOWED_NOTES:
        raise ValueError("Only PDF and DOCX lesson notes are supported.")
    path = UPLOAD_DIR / f"{secrets.token_hex(8)}-{filename}"
    file_storage.save(path)
    if extension == "pdf":
        if PdfReader is None:
            raise RuntimeError("Install pypdf to read PDF notes.")
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    else:
        if Document is None:
            raise RuntimeError("Install python-docx to read DOCX notes.")
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        text = "\n".join(paragraphs + tables)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 40:
        raise ValueError("The uploaded document did not contain readable text.")
    return text


def save_question_image(file_storage):
    if not file_storage or not file_storage.filename:
        return ""
    filename = secure_filename(file_storage.filename)
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in ALLOWED_IMAGES:
        raise ValueError("Question images must be PNG, JPG, JPEG, WEBP, or GIF.")
    stored_name = f"question-{secrets.token_hex(10)}.{extension}"
    file_storage.save(UPLOAD_DIR / stored_name)
    return stored_name


def class_group(class_name):
    value = (class_name or "").strip()
    match = re.match(r"([A-Za-z]+\s*\d+)", value)
    return match.group(1).replace(" ", "").upper() if match else value


def class_options():
    configured = [value.strip() for value in school_settings().get("managed_classes", "").split(",") if value.strip()]
    stored = [row["value"] for row in query("SELECT DISTINCT class_name AS value FROM users WHERE class_name != '' ORDER BY class_name")]
    values = configured + stored
    return sorted(set(values))


def class_group_options():
    return sorted(set(class_group(value) for value in class_options() if value))


def class_arm_options():
    return sorted(value for value in class_options() if re.search(r"[A-Za-z]\s*$", value))


def class_matches(value, selected):
    return not selected or value == selected or class_group(value) == class_group(selected)


def normalized_name(value):
    return " ".join(unicodedata.normalize("NFKC", value or "").casefold().split())


def valid_access_code(subject, code):
    record = query(
        "SELECT * FROM subject_access_codes WHERE subject = ? AND code = ? AND active = 1 ORDER BY id DESC",
        (subject.strip(), code.strip().upper()), one=True,
    )
    if not record:
        return False
    try:
        return datetime.now() <= datetime.fromisoformat(record["expires_at"])
    except ValueError:
        return False


def find_student_by_name(full_name):
    target = normalized_name(full_name)
    return next(
        (row for row in query("SELECT * FROM users WHERE role='student'") if normalized_name(row["full_name"]) == target),
        None,
    )


def generate_local_questions(subject, topic, notes, count, difficulty):
    sentences = [
        part.strip(" -:;")
        for part in re.split(r"(?<=[.!?])\s+|\n+", notes)
        if len(part.strip()) >= 25
    ]
    sentences = list(dict.fromkeys(sentences))
    if not sentences:
        sentences = [f"Key concepts and foundational principles of {topic}."]
    stems = [
        "According to the lesson notes, which statement is correct?",
        f"Which idea about {topic} is supported by the lesson notes?",
        f"What is the best description of {topic} based on the lesson?",
        f"Which statement accurately applies the lesson on {topic}?",
        f"Which conclusion can be drawn from the notes about {topic}?",
    ]
    if difficulty == "easy":
        stems = [
            "Which statement is directly stated in the lesson notes?",
            f"Which basic fact about {topic} appears in the lesson?",
        ]
    elif difficulty == "hard":
        stems = [
            f"Which conclusion about {topic} is best supported by the lesson?",
            f"Which application of the lesson on {topic} is most accurate?",
            f"Based on the notes, which interpretation of {topic} is most defensible?",
        ]
    questions = []
    for index in range(count):
        answer = sentences[index % len(sentences)]
        distractors = [
            sentences[(index + offset) % len(sentences)]
            for offset in range(1, min(4, len(sentences)))
        ]
        distractors += [
            f"{topic} is unrelated to the ideas in the lesson.",
            f"The lesson says that {topic} should never be studied.",
            f"{topic} only applies outside the classroom.",
        ]
        options = [answer] + list(dict.fromkeys(distractors))[:3]
        while len(options) < 4:
            options.append(f"This statement is not supported by the notes about {topic}.")
        questions.append(
            {
                "subject": subject, "topic": topic,
                "stem": f"{stems[index % len(stems)]} (Question {index + 1})",
                "option_a": options[0], "option_b": options[1],
                "option_c": options[2], "option_d": options[3],
                "correct_option": "A",
                "difficulty": difficulty,
                "explanation": "The correct answer is taken directly from the uploaded lesson notes.",
                "standard": "Teacher-defined curriculum standard",
            }
        )
    return questions


def generate_questions(subject, topic, notes, count, difficulty):
    if not notes or len(notes.strip()) < 40:
        raise ValueError("The lesson notes were empty or too short to generate grounded questions.")
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return generate_local_questions(subject, topic, notes, count, difficulty), "local (GEMINI_API_KEY is not set)"
    if genai is None:
        return generate_local_questions(subject, topic, notes, count, difficulty), "local (google-genai is not installed)"
    prompt = f"""Create exactly {count} curriculum-aligned MCQs for the subject {subject}.
Lesson topic: {topic}
Lesson notes (use only this source; do not invent facts):
{notes[:50000]}
Requested difficulty: {difficulty}
Return only valid JSON array items with subject, topic, stem, option_a, option_b,
option_c, option_d, correct_option (A/B/C/D), explanation, and standard.
Use different concepts from the notes for different questions. Make every stem
meaningfully different, use plausible distractors, and clear age-appropriate wording."""
    try:
        model_name = os.getenv("GEMINI_MODEL", "gemini-3.6-flash")
        with genai.Client(api_key=api_key) as client:
            response = client.models.generate_content(
                model=model_name, contents=prompt
            )
        text = response.text.strip().replace("```json", "").replace("```", "").strip()
        generated = json.loads(text)
        if not isinstance(generated, list) or len(generated) != count:
            raise ValueError("Gemini returned an incomplete question set.")
        required = {"stem", "option_a", "option_b", "option_c", "option_d", "correct_option"}
        if any(not required.issubset(item) for item in generated):
            raise ValueError("Gemini returned an invalid question structure.")
        for item in generated:
            item["subject"] = subject
            item["topic"] = topic
            item["difficulty"] = difficulty
        return generated, "Gemini"
    except Exception as error:
        reason = str(error).replace("\n", " ")[:180]
        return generate_local_questions(subject, topic, notes, count, difficulty), f"local fallback (Gemini error: {reason})"


def save_questions(questions, user_id):
    for item in questions:
        execute(
            """INSERT INTO questions
            (subject, topic, class_name, image_path, stem, option_a, option_b, option_c, option_d,
             correct_option, difficulty, explanation, standard, created_by, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                item.get("subject", "General"), item["topic"], item.get("class_name", ""),
                item.get("image_path", ""), item["stem"],
                item["option_a"], item["option_b"], item["option_c"], item["option_d"],
                item.get("correct_option", "A").upper(), item.get("difficulty", "medium"),
                item.get("explanation", ""),
                item.get("standard", ""), user_id, datetime.now().isoformat(timespec="seconds"),
            ),
        )


def ranking_tables():
    rows = query(
        """SELECT a.id, a.score, a.total, e.subject, e.exam_type, u.full_name, u.class_name
        FROM attempts a JOIN exams e ON e.id = a.exam_id
        JOIN users u ON u.id = a.student_id WHERE u.role = 'student' AND a.submitted_at != ''"""
    )
    class_filter = request.args.get("class_name", "").strip()
    subject_filter = request.args.get("subject", "").strip()
    exam_type_filter = request.args.get("exam_type", "").strip()
    ranking_scope = request.args.get("scope", "all")
    rows = [
        {**dict(row), "class_group": class_group(row["class_name"])}
        for row in rows
    ]
    user = current_user()
    if user and user["role"] == "teacher":
        allowed = teacher_subjects(user)
        rows = [row for row in rows if row["subject"] in allowed]
    if class_filter:
        rows = [row for row in rows if row["class_name"] == class_filter or row["class_group"] == class_filter]
    if subject_filter:
        rows = [row for row in rows if row["subject"] == subject_filter]
    if exam_type_filter in EXAM_TYPES:
        rows = [row for row in rows if row["exam_type"] == exam_type_filter]
    if not rows:
        return [], []
    data = pd.DataFrame([dict(row) for row in rows])
    if class_filter and any(row["class_name"] != class_filter for row in rows):
        data["ranking_class"] = data["class_group"]
    else:
        data["ranking_class"] = data["class_name"]
    data["percentage"] = (data["score"] / data["total"] * 100).round(1)
    subject = data.groupby(["ranking_class", "subject", "full_name"], as_index=False).agg(
        score=("score", "sum"), total=("total", "sum")
    )
    subject = subject.rename(columns={"ranking_class": "class_name"})
    subject["percentage"] = (subject["score"] / subject["total"] * 100).round(1)
    subject["rank"] = subject.groupby(["class_name", "subject"])["percentage"].rank(
        method="min", ascending=False
    ).astype(int)
    subject = subject.sort_values(["class_name", "subject", "rank", "full_name"])
    overall = data.groupby(["ranking_class", "full_name"], as_index=False).agg(
        score=("score", "sum"), total=("total", "sum")
    )
    overall = overall.rename(columns={"ranking_class": "class_name"})
    overall["percentage"] = (overall["score"] / overall["total"] * 100).round(1)
    overall["rank"] = overall.groupby("class_name")["percentage"].rank(
        method="min", ascending=False
    ).astype(int)
    overall = overall.sort_values(["class_name", "rank", "full_name"])
    return (
        [] if ranking_scope == "overall" else subject.to_dict("records"),
        [] if ranking_scope == "subject" else overall.to_dict("records"),
    )


@app.context_processor
def inject_globals():
    settings = school_settings()
    return {"user": current_user(), "school": settings, "themes": sorted(THEMES)}


@app.route("/")
def index():
    return redirect(url_for("dashboard") if current_user() else url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = query("SELECT * FROM users WHERE username = ? AND active = 1", (request.form["username"].strip(),), one=True)
        if user and check_password(request.form["password"], user["password_hash"]):
            session["user_id"] = user["id"]
            return redirect(url_for("dashboard"))
        flash("Invalid username or password.", "error")
    return render_template("login.html")


@app.route("/register/teacher", methods=["GET", "POST"])
def register_teacher():
    if request.method == "POST":
        full_name = request.form["full_name"].strip()
        class_name = request.form["class_name"].strip()
        subjects = request.form["subjects"].strip()
        access_subject = request.form["access_subject"].strip()
        access_code = request.form["access_code"].strip()
        if not full_name or not class_name or not subjects or not valid_access_code(access_subject, access_code):
            flash("Name, class, subjects, and a valid active subject access code are required.", "error")
        else:
            username = "teacher." + secrets.token_hex(5)
            user_id = execute(
                "INSERT INTO users(username, full_name, password_hash, role, class_name, subjects, created_at) VALUES (?, ?, ?, 'teacher', ?, ?, ?)",
                (username, full_name, hash_password(secrets.token_urlsafe(18)), class_name, subjects, datetime.now().isoformat(timespec="seconds")),
            )
            session["user_id"] = user_id
            flash("Teacher registration completed.", "success")
            return redirect(url_for("dashboard"))
    return render_template("register_teacher.html", classes=class_group_options(), subjects=category_values("managed_subjects", "SELECT DISTINCT subject AS value FROM questions ORDER BY subject"))


@app.route("/register/student", methods=["GET", "POST"])
def register_student():
    classes = class_arm_options()
    subjects = category_values("managed_subjects", "SELECT DISTINCT subject AS value FROM questions ORDER BY subject")
    if request.method == "POST":
        full_name = request.form["full_name"].strip()
        class_name = request.form["class_name"].strip()
        access_subject = request.form["access_subject"].strip()
        access_code = request.form["access_code"].strip()
        if not full_name or class_name not in classes or not valid_access_code(access_subject, access_code):
            flash("Enter your name, valid class, and active subject access code.", "error")
        else:
            existing = find_student_by_name(full_name)
            if existing:
                execute("UPDATE users SET full_name=?, class_name=?, active=1 WHERE id=?", (full_name, class_name, existing["id"]))
                user_id = existing["id"]
            else:
                username = "student." + secrets.token_hex(5)
                user_id = execute(
                    "INSERT INTO users(username, full_name, password_hash, role, class_name, created_at) VALUES (?, ?, ?, 'student', ?, ?)",
                    (username, full_name, hash_password(secrets.token_urlsafe(18)), class_name, datetime.now().isoformat(timespec="seconds")),
                )
            session["user_id"] = user_id
            flash("Student record updated. Continue to your examinations.", "success")
            return redirect(url_for("dashboard"))
    return render_template("register_student.html", classes=classes, subjects=subjects)


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.get("/dashboard")
@login_required()
def dashboard():
    user = current_user()
    stats = {
        "questions": query("SELECT COUNT(*) AS count FROM questions", one=True)["count"],
        "exams": query("SELECT COUNT(*) AS count FROM exams", one=True)["count"],
        "attempts": query("SELECT COUNT(*) AS count FROM attempts", one=True)["count"],
    }
    subject_filter = request.args.get("subject", "").strip()
    exam_type_filter = request.args.get("exam_type", "").strip()
    status_filter = request.args.get("status", "").strip()
    exam_sql = "SELECT e.*, COUNT(eq.question_id) AS question_count FROM exams e LEFT JOIN exam_questions eq ON eq.exam_id=e.id WHERE 1=1"
    exam_params = []
    if subject_filter:
        exam_sql += " AND e.subject = ?"
        exam_params.append(subject_filter)
    if exam_type_filter in EXAM_TYPES:
        exam_sql += " AND e.exam_type = ?"
        exam_params.append(exam_type_filter)
    if status_filter in ("published", "draft"):
        exam_sql += " AND e.published = ?"
        exam_params.append(int(status_filter == "published"))
    if user["role"] == "student":
        exam_sql += " AND (e.class_name = '' OR e.class_name = ? OR e.class_name = ?)"
        exam_params.extend([user["class_name"], class_group(user["class_name"])])
    elif user["role"] == "teacher":
        subjects = teacher_subjects(user)
        if subjects:
            exam_sql += " AND e.subject IN (" + ",".join("?" for _ in subjects) + ")"
            exam_params.extend(subjects)
        else:
            exam_sql += " AND e.created_by = ?"
            exam_params.append(user["id"])
    exam_sql += " GROUP BY e.id ORDER BY e.id DESC"
    exams = query(exam_sql, exam_params)
    subjects = [row["subject"] for row in query("SELECT DISTINCT subject FROM exams ORDER BY subject")]
    classes = class_options()
    return render_template("dashboard.html", stats=stats, exams=exams, user=user, subjects=subjects, classes=classes, subject_filter=subject_filter, status_filter=status_filter, exam_type_filter=exam_type_filter, exam_types=EXAM_TYPES)


@app.route("/questions", methods=["GET", "POST"])
@login_required("admin", "teacher")
def questions():
    if request.method == "POST":
        try:
            subject = request.form["subject"].strip()
            if current_user()["role"] == "teacher" and subject not in teacher_subjects():
                raise ValueError("You can only create questions for your assigned subjects.")
            topic = request.form["topic"].strip()
            notes = request.form.get("notes", "").strip()
            if request.files.get("notes_file") and request.files["notes_file"].filename:
                notes = extract_notes(request.files["notes_file"])
            if not subject or not topic or not notes:
                raise ValueError("Subject, lesson topic, and notes are required.")
            difficulty = request.form.get("difficulty", "medium")
            generated, source = generate_questions(subject, topic, notes, int(request.form["count"]), difficulty)
            for item in generated:
                item["class_name"] = request.form.get("class_name", "").strip()
            save_questions(generated, current_user()["id"])
            flash(f"Read {len(notes):,} characters from the lesson source and saved {len(generated)} questions using {source}.", "success")
        except (ValueError, RuntimeError, json.JSONDecodeError) as error:
            flash(str(error), "error")
        except Exception as error:
            flash(f"Question generation failed: {error}", "error")
        return redirect(url_for("questions"))
    subject_filter = request.args.get("subject", "").strip()
    class_filter = request.args.get("class_name", "").strip()
    sql = "SELECT * FROM questions WHERE 1=1"
    params = []
    if subject_filter:
        sql += " AND subject = ?"
        params.append(subject_filter)
    scope_sql, scope_params = scoped_question_sql(current_user())
    sql += scope_sql
    params.extend(scope_params)
    sql += " ORDER BY id DESC"
    return render_template(
        "questions.html", questions=[row for row in query(sql, params) if class_matches(row["class_name"], class_filter)],
        subjects=teacher_subjects() if current_user()["role"] == "teacher" else category_values("managed_subjects", "SELECT DISTINCT subject AS value FROM questions ORDER BY subject"),
        classes=class_group_options(),
        subject_filter=subject_filter, class_filter=class_filter,
        teacher_subjects=teacher_subjects(),
    )


@app.post("/questions/clear")
@login_required("admin", "teacher")
def clear_questions():
    user = current_user()
    if user["role"] == "admin":
        execute("DELETE FROM questions")
    else:
        subjects = teacher_subjects(user)
        if subjects:
            placeholders = ",".join("?" for _ in subjects)
            execute(f"DELETE FROM questions WHERE created_by = ? OR subject IN ({placeholders})", [user["id"], *subjects])
        else:
            execute("DELETE FROM questions WHERE created_by = ?", (user["id"],))
    flash("Saved questions cleared.", "success")
    return redirect(url_for("questions"))


@app.route("/questions/manual", methods=["GET", "POST"])
@login_required("admin", "teacher")
def manual_question():
    if request.method == "POST":
        if current_user()["role"] == "teacher" and request.form.get("subject", "").strip() not in teacher_subjects():
            flash("You can only create questions for your assigned subjects.", "error")
            return redirect(url_for("manual_question"))
        stems = request.form.getlist("stem[]") or [request.form.get("stem", "")]
        option_a = request.form.getlist("option_a[]") or [request.form.get("option_a", "")]
        option_b = request.form.getlist("option_b[]") or [request.form.get("option_b", "")]
        option_c = request.form.getlist("option_c[]") or [request.form.get("option_c", "")]
        option_d = request.form.getlist("option_d[]") or [request.form.get("option_d", "")]
        correct_options = request.form.getlist("correct_option[]") or [request.form.get("correct_option", "")]
        images = request.files.getlist("question_image[]") or [request.files.get("question_image")]
        fields = [stems, option_a, option_b, option_c, option_d, correct_options]
        if not all(len(field) == len(stems) for field in fields):
            flash("Each manual question must include all four answer choices and a correct option.", "error")
            return redirect(url_for("manual_question"))
        if not stems or not all(stem.strip() for stem in stems):
            flash("Add at least one complete manual question.", "error")
            return redirect(url_for("manual_question"))
        try:
            for index, stem in enumerate(stems):
                image_path = save_question_image(images[index] if index < len(images) else None)
                values = (
                    request.form["subject"].strip(), request.form["topic"].strip(),
                    request.form.get("class_name", "").strip(), image_path, stem.strip(),
                    option_a[index].strip(), option_b[index].strip(),
                    option_c[index].strip(), option_d[index].strip(),
                    correct_options[index], request.form["difficulty"],
                    request.form.get("explanation", "").strip(), request.form.get("standard", "").strip(),
                    current_user()["id"], datetime.now().isoformat(timespec="seconds"),
                )
                if not all((values[0], values[1], values[2], values[4], values[5], values[6], values[7], values[8], values[9])):
                    raise ValueError("Subject, topic, question, all answers, and correct option are required.")
                execute(
                """INSERT INTO questions
                (subject, topic, class_name, image_path, stem, option_a, option_b, option_c, option_d,
                 correct_option, difficulty, explanation, standard, created_by, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                values,
                )
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("manual_question"))
        flash(f"Added {len(stems)} manual question{'s' if len(stems) != 1 else ''}.", "success")
        return redirect(url_for("questions"))
    return render_template(
        "question_manual.html",
        subjects=teacher_subjects() if current_user()["role"] == "teacher" else category_values("managed_subjects", "SELECT DISTINCT subject AS value FROM questions ORDER BY subject"),
        classes=class_group_options(),
    )


@app.route("/questions/<int:question_id>/edit", methods=["GET", "POST"])
@login_required("admin", "teacher")
def edit_question(question_id):
    question = query("SELECT * FROM questions WHERE id = ?", (question_id,), one=True)
    if not question:
        flash("Question not found.", "error")
        return redirect(url_for("questions"))
    if current_user()["role"] == "teacher" and question["subject"] not in teacher_subjects():
        flash("Teachers can edit only questions in their assigned subjects.", "error")
        return redirect(url_for("questions"))
    if request.method == "POST":
        try:
            uploaded_image = save_question_image(request.files.get("question_image"))
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("edit_question", question_id=question_id))
        image_path = uploaded_image or question["image_path"]
        values = (
            request.form["subject"].strip(), request.form["topic"].strip(),
            request.form.get("class_name", "").strip(), image_path, request.form["stem"].strip(),
            request.form["option_a"].strip(),
            request.form["option_b"].strip(), request.form["option_c"].strip(),
            request.form["option_d"].strip(), request.form["correct_option"],
            request.form["difficulty"], request.form["explanation"].strip(),
            request.form["standard"].strip(),
            question_id,
        )
        if not all(values[:9]):
            flash("Subject, topic, question, all answers, and correct option are required.", "error")
        else:
            execute(
                """UPDATE questions SET subject=?, topic=?, class_name=?, image_path=?, stem=?, option_a=?, option_b=?,
                option_c=?, option_d=?, correct_option=?, difficulty=?, explanation=?, standard=? WHERE id=?""",
                values,
            )
            flash("Question updated.", "success")
            return redirect(url_for("questions"))
    return render_template("question_edit.html", question=question)


@app.post("/questions/<int:question_id>/delete")
@login_required("admin", "teacher")
def delete_question(question_id):
    question = query("SELECT created_by FROM questions WHERE id = ?", (question_id,), one=True)
    user = current_user()
    if not question or (user["role"] == "teacher" and query("SELECT subject FROM questions WHERE id=?", (question_id,), one=True)["subject"] not in teacher_subjects(user)):
        flash("Teachers can delete only questions in their assigned subjects.", "error")
        return redirect(url_for("questions"))
    execute("DELETE FROM questions WHERE id = ?", (question_id,))
    flash("Question deleted.", "success")
    return redirect(url_for("questions"))


@app.post("/admin/users/<int:user_id>/delete")
@login_required("admin")
def delete_user(user_id):
    account = query("SELECT username, role FROM users WHERE id = ?", (user_id,), one=True)
    if not account or account["role"] == "admin":
        flash("Only teacher and student accounts can be deleted here.", "error")
    elif user_id == current_user()["id"]:
        flash("You cannot delete your own account.", "error")
    else:
        execute("UPDATE users SET active = 0 WHERE id = ?", (user_id,))
        flash(f"{account['role'].capitalize()} access removed.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/users/<int:user_id>/status")
@login_required("admin")
def update_user_status(user_id):
    account = query("SELECT role, active FROM users WHERE id = ?", (user_id,), one=True)
    if not account or account["role"] == "admin":
        flash("Only teacher and student accounts can be managed here.", "error")
    else:
        execute("UPDATE users SET active = ? WHERE id = ?", (int(request.form.get("active") == "1"), user_id))
        flash("User access updated.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/users/<int:user_id>/remove")
@login_required("admin")
def remove_user(user_id):
    account = query("SELECT role FROM users WHERE id = ?", (user_id,), one=True)
    if not account or account["role"] == "admin" or user_id == current_user()["id"]:
        flash("This administrator account cannot be permanently removed here.", "error")
    else:
        execute("DELETE FROM responses WHERE attempt_id IN (SELECT id FROM attempts WHERE student_id = ?)", (user_id,))
        execute("DELETE FROM rewrite_permissions WHERE attempt_id IN (SELECT id FROM attempts WHERE student_id = ?)", (user_id,))
        execute("DELETE FROM attempts WHERE student_id = ?", (user_id,))
        execute("DELETE FROM users WHERE id = ?", (user_id,))
        flash("User permanently removed.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/exams/<int:exam_id>/publish")
@login_required("admin")
def publish_exam(exam_id):
    exam = query("SELECT published, title FROM exams WHERE id = ?", (exam_id,), one=True)
    if not exam:
        flash("Exam not found.", "error")
    else:
        execute("UPDATE exams SET published = ? WHERE id = ?", (int(not exam["published"]), exam_id))
        flash(f"{exam['title']} is now {'published' if not exam['published'] else 'unpublished'}.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/exams/<int:exam_id>/remove")
@login_required("admin")
def remove_exam(exam_id):
    exam = query("SELECT title FROM exams WHERE id = ?", (exam_id,), one=True)
    if not exam:
        flash("Exam not found.", "error")
    else:
        execute("DELETE FROM responses WHERE attempt_id IN (SELECT id FROM attempts WHERE exam_id = ?)", (exam_id,))
        execute("DELETE FROM rewrite_permissions WHERE attempt_id IN (SELECT id FROM attempts WHERE exam_id = ?)", (exam_id,))
        execute("DELETE FROM attempts WHERE exam_id = ?", (exam_id,))
        execute("DELETE FROM exam_questions WHERE exam_id = ?", (exam_id,))
        execute("DELETE FROM exams WHERE id = ?", (exam_id,))
        flash(f"{exam['title']} was permanently removed.", "success")
    return redirect(url_for("admin"))


@app.route("/exams/new", methods=["GET", "POST"])
@login_required("admin", "teacher")
def new_exam():
    subject_filter = request.args.get("subject", "").strip()
    class_filter = request.args.get("class_name", "").strip()
    question_sql = "SELECT * FROM questions WHERE 1=1"
    question_params = []
    if subject_filter:
        question_sql += " AND subject = ?"
        question_params.append(subject_filter)
    scope_sql, scope_params = scoped_question_sql(current_user())
    question_sql += scope_sql
    question_params.extend(scope_params)
    question_sql += " ORDER BY subject, topic, id"
    available = [row for row in query(question_sql, question_params) if class_matches(row["class_name"], class_filter)]
    if request.method == "POST":
        selected = request.form.getlist("question_ids")
        if not request.form["title"].strip() or not selected:
            flash("Provide an exam title and select at least one question.", "error")
        else:
            exam_id = execute(
                """INSERT INTO exams(title, subject, class_name, exam_type, instructions, duration_minutes, created_by, published, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    request.form["title"].strip(), request.form["subject"].strip(),
                    request.form["class_name"].strip(), request.form["exam_type"], request.form["instructions"], int(request.form["duration_minutes"]),
                    current_user()["id"], int("published" in request.form),
                    datetime.now().isoformat(timespec="seconds"),
                ),
            )
            for position, question_id in enumerate(selected, 1):
                execute("INSERT INTO exam_questions(exam_id, question_id, position) VALUES (?, ?, ?)", (exam_id, int(question_id), position))
            flash("Exam created.", "success")
            return redirect(url_for("dashboard"))
    return render_template(
        "exam_new.html", questions=available, exam_types=EXAM_TYPES,
        subjects=category_values("managed_subjects", "SELECT DISTINCT subject AS value FROM questions ORDER BY subject"),
        classes=class_group_options(),
        subject_filter=subject_filter, class_filter=class_filter,
    )


@app.route("/exams/<int:exam_id>/take", methods=["GET", "POST"])
@login_required("student")
def take_exam(exam_id):
    exam = query("SELECT * FROM exams WHERE id = ? AND published = 1", (exam_id,), one=True)
    if not exam:
        flash("Exam is not available.", "error")
        return redirect(url_for("dashboard"))
    existing = query(
        "SELECT id, started_at, submitted_at FROM attempts WHERE exam_id = ? AND student_id = ?",
        (exam_id, current_user()["id"]), one=True,
    )
    permission = query(
        "SELECT id FROM rewrite_permissions WHERE attempt_id = ?",
        (existing["id"],), one=True,
    ) if existing else None
    if existing and existing["submitted_at"] and not permission:
        flash("You have already submitted this exam.", "error")
        return redirect(url_for("dashboard"))
    force_submitted = False
    if existing and existing["submitted_at"] and permission:
        # Reuse the attempt row with a fresh timer. This avoids a unique-key
        # collision if the student opens the rewrite link more than once.
        now = datetime.now().isoformat(timespec="seconds")
        execute("DELETE FROM responses WHERE attempt_id = ?", (existing["id"],))
        execute("DELETE FROM rewrite_permissions WHERE id = ?", (permission["id"],))
        execute(
            "UPDATE attempts SET score = 0, total = 0, started_at = ?, submitted_at = '' WHERE id = ?",
            (now, existing["id"]),
        )
        existing = query(
            "SELECT id, started_at, submitted_at FROM attempts WHERE id = ?",
            (existing["id"],), one=True,
        )
        permission = None
    if existing and existing["started_at"]:
        try:
            started = datetime.fromisoformat(existing["started_at"])
        except ValueError:
            started = datetime.now()
            execute(
                "UPDATE attempts SET started_at = ?, submitted_at = '' WHERE id = ?",
                (started.isoformat(timespec="seconds"), existing["id"]),
            )
        if (datetime.now() - started).total_seconds() > exam["duration_minutes"] * 60:
            if permission:
                now = datetime.now().isoformat(timespec="seconds")
                execute("DELETE FROM responses WHERE attempt_id = ?", (existing["id"],))
                execute("DELETE FROM rewrite_permissions WHERE id = ?", (permission["id"],))
                execute("UPDATE attempts SET score=0, total=0, started_at=?, submitted_at='' WHERE id=?", (now, existing["id"]))
                existing = query("SELECT id, started_at, submitted_at FROM attempts WHERE id=?", (existing["id"],), one=True)
                permission = None
            elif request.method == "POST":
                force_submitted = True
            else:
                flash("This exam window has expired.", "error")
                return redirect(url_for("dashboard"))
    questions = query("SELECT q.* FROM questions q JOIN exam_questions eq ON eq.question_id=q.id WHERE eq.exam_id=? ORDER BY eq.position", (exam_id,))
    if request.method == "POST":
        answers = {question["id"]: request.form.get(f"question_{question['id']}", "") for question in questions}
        score = sum(answer == question["correct_option"] for question, answer in zip(questions, answers.values()))
        started_at = existing and query("SELECT started_at FROM attempts WHERE id = ?", (existing["id"],), one=True)["started_at"] or datetime.now().isoformat(timespec="seconds")
        now = datetime.now().isoformat(timespec="seconds")
        if existing:
            execute("DELETE FROM responses WHERE attempt_id = ?", (existing["id"],))
            execute(
                "UPDATE attempts SET score = ?, total = ?, started_at = ?, submitted_at = ? WHERE id = ?",
                (score, len(questions), started_at, now, existing["id"]),
            )
            attempt_id = existing["id"]
        else:
            attempt_id = execute("INSERT INTO attempts(exam_id, student_id, score, total, started_at, submitted_at) VALUES (?, ?, ?, ?, ?, ?)", (exam_id, current_user()["id"], score, len(questions), started_at, now))
        for question in questions:
            execute("INSERT INTO responses(attempt_id, question_id, answer, correct) VALUES (?, ?, ?, ?)", (attempt_id, question["id"], answers[question["id"]], int(answers[question["id"]] == question["correct_option"])))
        return render_template("result.html", exam=exam, force_submitted=force_submitted)
    if not existing:
        now = datetime.now().isoformat(timespec="seconds")
        execute(
            "INSERT INTO attempts(exam_id, student_id, score, total, started_at, submitted_at) VALUES (?, ?, 0, ?, ?, '')",
            (exam_id, current_user()["id"], len(questions), now),
        )
    started_attempt = query("SELECT started_at FROM attempts WHERE exam_id = ? AND student_id = ?", (exam_id, current_user()["id"]), one=True)
    elapsed = max(0, int((datetime.now() - datetime.fromisoformat(started_attempt["started_at"])).total_seconds()))
    remaining_seconds = max(0, exam["duration_minutes"] * 60 - elapsed)
    return render_template("take_exam.html", exam=exam, questions=questions, remaining_seconds=remaining_seconds)


@app.post("/admin/rewrite/<int:attempt_id>")
@login_required("admin")
def grant_rewrite(attempt_id):
    attempt = query(
        "SELECT a.id, u.full_name, e.title FROM attempts a JOIN users u ON u.id=a.student_id JOIN exams e ON e.id=a.exam_id WHERE a.id=?",
        (attempt_id,), one=True
    )
    if not attempt:
        flash("Attempt not found.", "error")
    else:
        execute(
            "INSERT OR REPLACE INTO rewrite_permissions(attempt_id, granted_by, granted_at) VALUES (?, ?, ?)",
            (attempt_id, current_user()["id"], datetime.now().isoformat(timespec="seconds")),
        )
        flash(f"Rewrite access granted to {attempt['full_name']} for {attempt['title']}.", "success")
    return redirect(url_for("admin"))


@app.route("/rankings")
@login_required("admin", "teacher")
def rankings():
    subject_rows, overall_rows = ranking_tables()
    class_values = [row["class_name"] for row in query("SELECT DISTINCT class_name FROM users WHERE role='student' AND class_name != ''")]
    return render_template(
        "rankings.html", subject_rows=subject_rows, overall_rows=overall_rows,
        classes=sorted(set(class_values + [class_group(value) for value in class_values])),
        subjects=[row["subject"] for row in query("SELECT DISTINCT subject FROM exams ORDER BY subject")],
        exam_types=EXAM_TYPES, exam_type_filter=request.args.get("exam_type", ""),
        ranking_scope=request.args.get("scope", "all"),
        class_filter=request.args.get("class_name", ""), subject_filter=request.args.get("subject", ""),
    )


@app.route("/analytics")
@login_required("admin", "teacher")
def analytics():
    class_filter = request.args.get("class_name", "").strip()
    subject_filter = request.args.get("subject", "").strip()
    exam_type_filter = request.args.get("exam_type", "").strip()
    rows = query(
        """SELECT u.full_name, u.class_name, e.subject, e.exam_type,
        a.score, a.total
        FROM attempts a
        JOIN users u ON u.id = a.student_id
        JOIN exams e ON e.id = a.exam_id
        WHERE u.role = 'student' AND a.submitted_at != ''"""
    )
    records = [dict(row) for row in rows]
    user = current_user()
    if user["role"] == "teacher":
        records = [row for row in records if row["subject"] in teacher_subjects(user)]
    if class_filter:
        records = [row for row in records if class_matches(row["class_name"], class_filter)]
    if subject_filter:
        records = [row for row in records if row["subject"] == subject_filter]
    if exam_type_filter in EXAM_TYPES:
        records = [row for row in records if row["exam_type"] == exam_type_filter]
    frame = pd.DataFrame(records)
    if frame.empty:
        class_chart, subject_chart, type_chart, distribution = [], [], [], []
        stats = {"students": 0, "submissions": 0, "average": 0, "pass_rate": 0}
    else:
        frame["percentage"] = (frame["score"] / frame["total"] * 100).round(1)
        class_chart = frame.groupby(frame["class_name"].map(class_group), as_index=False).agg(
            average=("percentage", "mean"), submissions=("percentage", "count")
        ).round(1).to_dict("records")
        class_chart = [{"label": row["class_name"], "average": row["average"], "submissions": row["submissions"]} for row in class_chart]
        subject_chart = frame.groupby("subject", as_index=False).agg(
            average=("percentage", "mean"), submissions=("percentage", "count")
        ).round(1).to_dict("records")
        type_chart = frame.groupby("exam_type", as_index=False).agg(
            average=("percentage", "mean"), submissions=("percentage", "count")
        ).round(1).to_dict("records")
        bins = pd.cut(frame["percentage"], bins=[-1, 39, 49, 59, 69, 100], labels=["0-39", "40-49", "50-59", "60-69", "70-100"])
        distribution = [{"label": str(label), "count": int((bins == label).sum())} for label in bins.cat.categories]
        stats = {
            "students": int(frame["full_name"].nunique()),
            "submissions": int(len(frame)),
            "average": round(float(frame["percentage"].mean()), 1),
            "pass_rate": round(float((frame["percentage"] >= 50).mean() * 100), 1),
        }
    return render_template(
        "analytics.html",
        class_chart=class_chart, subject_chart=subject_chart,
        type_chart=type_chart, distribution=distribution, stats=stats,
        classes=sorted(set([row["class_name"] for row in query("SELECT DISTINCT class_name FROM users WHERE role='student' AND class_name != ''")] + class_group_options())),
        subjects=teacher_subjects(user) if user["role"] == "teacher" else category_values("managed_subjects", "SELECT DISTINCT subject AS value FROM questions ORDER BY subject"),
        exam_types=EXAM_TYPES, class_filter=class_filter,
        subject_filter=subject_filter, exam_type_filter=exam_type_filter,
    )


@app.get("/reports/results.csv")
@login_required("admin", "teacher")
def results_csv():
    class_filter = request.args.get("class_name", "").strip()
    exam_type_filter = request.args.get("exam_type", "").strip()
    rows = query(
        """SELECT u.full_name, u.username, u.class_name, e.title, e.subject, e.exam_type,
        a.score, a.total, a.submitted_at
        FROM attempts a
        JOIN users u ON u.id = a.student_id
        JOIN exams e ON e.id = a.exam_id
        WHERE u.role = 'student' AND a.submitted_at != ''
        ORDER BY u.class_name, u.full_name, e.exam_type, e.title"""
    )
    records = [dict(row) for row in rows]
    if current_user()["role"] == "teacher":
        records = [row for row in records if row["subject"] in teacher_subjects()]
    if class_filter:
        records = [row for row in records if row["class_name"] == class_filter or class_group(row["class_name"]) == class_filter]
        if any(row["class_name"] != class_filter for row in records):
            for row in records:
                row["class_name"] = class_filter
    if exam_type_filter in EXAM_TYPES:
        records = [row for row in records if row["exam_type"] == exam_type_filter]
    frame = pd.DataFrame(records)
    if frame.empty:
        frame = pd.DataFrame(columns=["class_name", "full_name", "username", "overall_score", "overall_total", "overall_percentage"])
    else:
        summary = frame.groupby(["class_name", "full_name", "username"], as_index=False).agg(
            overall_score=("score", "sum"), overall_total=("total", "sum")
        )
        summary["overall_percentage"] = (
            summary["overall_score"] / summary["overall_total"] * 100
        ).round(1)
        for exam_type in sorted(frame["exam_type"].dropna().unique()):
            type_frame = frame[frame["exam_type"] == exam_type]
            type_summary = type_frame.groupby(
                ["class_name", "full_name", "username"], as_index=False
            ).agg(score=("score", "sum"), total=("total", "sum"))
            key = re.sub(r"[^a-z0-9]+", "_", exam_type.lower()).strip("_")
            type_summary[f"{key}_score"] = type_summary["score"]
            type_summary[f"{key}_total"] = type_summary["total"]
            type_summary[f"{key}_percentage"] = (
                type_summary["score"] / type_summary["total"] * 100
            ).round(1)
            summary = summary.merge(
                type_summary.drop(columns=["score", "total"]),
                on=["class_name", "full_name", "username"], how="left",
            )
        frame = summary
    output = io.StringIO()
    frame.to_csv(output, index=False)
    output.seek(0)
    download = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    suffix = re.sub(r"[^a-z0-9]+", "-", class_filter.lower()).strip("-") or "all-classes"
    return send_file(download, mimetype="text/csv", as_attachment=True, download_name=f"student-results-{suffix}.csv")


@app.route("/admin", methods=["GET", "POST"])
@login_required("admin")
def admin():
    if request.method == "POST":
        action = request.form["action"]
        if action == "branding":
            logo = request.files.get("logo")
            logo_name = school_settings().get("school_logo", "")
            if logo and logo.filename:
                extension = Path(secure_filename(logo.filename)).suffix.lower().lstrip(".")
                if extension not in ALLOWED_LOGOS:
                    flash("Logo must be PNG, JPG, JPEG, or WEBP.", "error")
                    return redirect(url_for("admin"))
                logo_name = f"school-logo.{extension}"
                logo.save(UPLOAD_DIR / logo_name)
            theme = request.form.get("theme", "ocean")
            if theme not in THEMES:
                theme = "ocean"
            theme_image = school_settings().get("theme_image", "")
            theme_upload = request.files.get("theme_image")
            if theme_upload and theme_upload.filename:
                extension = Path(secure_filename(theme_upload.filename)).suffix.lower().lstrip(".")
                if extension not in ALLOWED_LOGOS:
                    flash("Theme image must be PNG, JPG, JPEG, or WEBP.", "error")
                    return redirect(url_for("admin"))
                theme_image = f"theme-image.{extension}"
                theme_upload.save(UPLOAD_DIR / theme_image)
                theme = "image"
            for key in ("school_name", "school_address"):
                execute("INSERT OR REPLACE INTO settings(key, value) VALUES (?, ?)", (key, request.form[key].strip()))
            execute("INSERT OR REPLACE INTO settings(key, value) VALUES (?, ?)", ("school_logo", logo_name))
            execute("INSERT OR REPLACE INTO settings(key, value) VALUES (?, ?)", ("theme", theme))
            execute("INSERT OR REPLACE INTO settings(key, value) VALUES (?, ?)", ("theme_image", theme_image))
            flash("School branding updated.", "success")
        elif action == "user":
            try:
                execute("INSERT INTO users(username, full_name, password_hash, role, class_name, subjects, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)", (request.form["username"], request.form["full_name"], hash_password(request.form["password"]), request.form["role"], request.form.get("class_name", "").strip(), request.form.get("subjects", "").strip(), datetime.now().isoformat(timespec="seconds")))
                flash("User created.", "success")
            except sqlite3.IntegrityError:
                flash("Username already exists.", "error")
        elif action == "students_csv":
            upload = request.files.get("students_csv")
            if not upload or not upload.filename.lower().endswith(".csv"):
                flash("Choose a CSV file containing student details.", "error")
            else:
                try:
                    frame = pd.read_csv(upload).fillna("")
                    required = {"full_name", "password", "class_name"}
                    if not required.issubset(frame.columns):
                        raise ValueError("CSV columns must include full_name, password, and class_name. username is optional.")
                    created = 0
                    for row in frame.to_dict("records"):
                        username = str(row.get("username", "")).strip()
                        full_name = str(row["full_name"]).strip()
                        class_name = str(row["class_name"]).strip()
                        existing = find_student_by_name(full_name)
                        if existing:
                            execute("UPDATE users SET full_name=?, class_name=?, active=1 WHERE id=?", (full_name, class_name, existing["id"]))
                            created += 1
                            continue
                        if not username:
                            username = re.sub(r"[^a-z0-9]+", ".", full_name.lower()).strip(".")
                        execute(
                            """INSERT INTO users(username, full_name, password_hash, role, class_name, created_at)
                            VALUES (?, ?, ?, 'student', ?, ?)""",
                            (username, full_name, hash_password(str(row["password"])), class_name, datetime.now().isoformat(timespec="seconds")),
                        )
                        created += 1
                    flash(f"Imported {created} students.", "success")
                except (ValueError, KeyError, sqlite3.IntegrityError) as error:
                    flash(f"Student CSV import failed: {error}", "error")
        elif action == "teachers_csv":
            upload = request.files.get("teachers_csv")
            if not upload or not upload.filename.lower().endswith(".csv"):
                flash("Choose a CSV file containing teacher details.", "error")
            else:
                try:
                    frame = pd.read_csv(upload).fillna("")
                    required = {"full_name", "password"}
                    if not required.issubset(frame.columns):
                        raise ValueError("Teacher CSV columns must include full_name and password. username and class_name are optional.")
                    created = 0
                    for row in frame.to_dict("records"):
                        full_name = str(row["full_name"]).strip()
                        username = str(row.get("username", "")).strip() or re.sub(r"[^a-z0-9]+", ".", full_name.lower()).strip(".")
                        execute(
                            """INSERT INTO users(username, full_name, password_hash, role, class_name, subjects, created_at)
                            VALUES (?, ?, ?, 'teacher', ?, ?, ?)""",
                            (username, full_name, hash_password(str(row["password"])), str(row.get("class_name", "")).strip(), str(row.get("subjects", "")).strip(), datetime.now().isoformat(timespec="seconds")),
                        )
                        created += 1
                    flash(f"Imported {created} teachers.", "success")
                except (ValueError, KeyError, sqlite3.IntegrityError) as error:
                    flash(f"Teacher CSV import failed: {error}", "error")
        elif action == "exam_settings":
            try:
                duration = int(request.form["duration_minutes"])
                if duration < 1:
                    raise ValueError("Timer must be at least one minute.")
                execute(
                    "UPDATE exams SET duration_minutes = ?, exam_type = ?, class_name = ? WHERE id = ?",
                    (duration, request.form["exam_type"], request.form["class_name"].strip(), int(request.form["exam_id"])),
                )
                flash("Exam timer and type updated.", "success")
            except (ValueError, KeyError):
                flash("Enter a valid exam timer and type.", "error")
        elif action == "categories":
            subjects = ",".join(sorted({value.strip() for value in request.form.get("subjects", "").split(",") if value.strip()}))
            classes = ",".join(sorted({value.strip() for value in request.form.get("classes", "").split(",") if value.strip()}))
            execute("INSERT OR REPLACE INTO settings(key, value) VALUES ('managed_subjects', ?)", (subjects,))
            execute("INSERT OR REPLACE INTO settings(key, value) VALUES ('managed_classes', ?)", (classes,))
            flash("Subject and class categories updated.", "success")
        elif action == "access_code":
            subject = request.form.get("access_subject", "").strip()
            try:
                duration_hours = float(request.form.get("duration_hours", "24"))
                if not subject or duration_hours <= 0 or duration_hours > 720:
                    raise ValueError
                code = secrets.token_urlsafe(8).replace("-", "").replace("_", "").upper()[:10]
                expires_at = datetime.now() + timedelta(hours=duration_hours)
                execute(
                    "INSERT INTO subject_access_codes(subject, code, expires_at, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
                    (subject, code, expires_at.isoformat(timespec="seconds"), current_user()["id"], datetime.now().isoformat(timespec="seconds")),
                )
                flash(f"Access code for {subject}: {code} (expires {expires_at.strftime('%Y-%m-%d %H:%M')}).", "success")
            except ValueError:
                flash("Enter a subject and a duration between 0 and 720 hours.", "error")
        elif action == "clear_data":
            if not secrets.compare_digest(request.form.get("clear_password", ""), DATA_CLEAR_PASSWORD):
                flash("The special data-clear password is incorrect.", "error")
            else:
                admin_id = current_user()["id"]
                connection = db()
                connection.executescript(
                    """
                    DELETE FROM responses;
                    DELETE FROM rewrite_permissions;
                    DELETE FROM attempts;
                    DELETE FROM exam_questions;
                    DELETE FROM exams;
                    DELETE FROM questions;
                    DELETE FROM subject_access_codes;
                    DELETE FROM users WHERE id != %d;
                    DELETE FROM settings;
                    """ % admin_id
                )
                connection.commit()
                connection.close()
                flash("All application data was cleared. The current administrator account was retained.", "success")
        return redirect(url_for("admin"))

    role_filter = request.args.get("role", "").strip()
    class_filter = request.args.get("class_name", "").strip()
    user_sql = "SELECT id, username, full_name, role, class_name, subjects, active FROM users WHERE role != 'admin'"
    user_params = []
    if role_filter in ("teacher", "student"):
        user_sql += " AND role = ?"
        user_params.append(role_filter)
    if class_filter:
        user_sql += " AND class_name = ?"
        user_params.append(class_filter)
    user_sql += " ORDER BY id"
    attempts = query(
        """SELECT a.id, a.score, a.total, a.submitted_at, u.full_name, u.class_name, e.title,
        EXISTS(SELECT 1 FROM rewrite_permissions rp WHERE rp.attempt_id=a.id) AS rewrite_granted
        FROM attempts a JOIN users u ON u.id=a.student_id JOIN exams e ON e.id=a.exam_id
        ORDER BY CASE WHEN a.submitted_at = '' THEN 0 ELSE 1 END, a.submitted_at DESC"""
    )
    return render_template(
        "admin.html",
        users=query(user_sql, user_params),
        attempts=attempts,
        role_filter=role_filter, class_filter=class_filter,
        classes=class_options(),
        exam_classes=class_group_options(),
        managed_subjects=category_values("managed_subjects", "SELECT DISTINCT subject AS value FROM questions ORDER BY subject"),
        managed_classes=class_options(),
        access_codes=query("SELECT * FROM subject_access_codes ORDER BY id DESC"),
        exams=query("SELECT id, title, subject, class_name, exam_type, duration_minutes, published FROM exams ORDER BY id DESC"),
        exam_types=EXAM_TYPES,
    )


@app.post("/admin/access-codes/<int:code_id>/toggle")
@login_required("admin")
def toggle_access_code(code_id):
    record = query("SELECT active, subject FROM subject_access_codes WHERE id = ?", (code_id,), one=True)
    if not record:
        flash("Access code not found.", "error")
    else:
        execute("UPDATE subject_access_codes SET active = ? WHERE id = ?", (int(not record["active"]), code_id))
        flash(f"{record['subject']} access code {'enabled' if not record['active'] else 'disabled'}.", "success")
    return redirect(url_for("admin"))


@app.post("/admin/access-codes/<int:code_id>/delete")
@login_required("admin")
def delete_access_code(code_id):
    record = query("SELECT subject FROM subject_access_codes WHERE id = ?", (code_id,), one=True)
    if not record:
        flash("Access code not found.", "error")
    else:
        execute("DELETE FROM subject_access_codes WHERE id = ?", (code_id,))
        flash(f"{record['subject']} access code was permanently deleted.", "success")
    return redirect(url_for("admin"))


@app.get("/uploads/<path:filename>")
def uploads(filename):
    return send_file(UPLOAD_DIR / secure_filename(filename))


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=os.getenv("FLASK_DEBUG") == "1")
